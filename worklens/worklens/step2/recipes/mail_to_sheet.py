"""レシピ: 問い合わせメールの内容を台帳へ入力する。

STEP1 が第2位に挙げた「メールで受けた内容をExcelへ入力」の実装。

**STEP1 との違いを明確にしておく。**
STEP1（業務の発見）はメール本文を一切読まない。読む必要が無いからで、
その約束は scopes.py に書いてある。
STEP2（自動化の実行）は、ユーザーがその業務を「自動化したい」と選んだ場合に限り、
業務としてメール本文を扱う。人がやっていた作業を代行するのだから当然だが、
**観測のために本文を読むことと、業務として本文を処理することは別物**である。
この境界を曖昧にしないために、本文はレシピの中だけで扱い、
収集データベースには一切保存しない。

抽出は2系統:
  - LLM（Claude）: 自由文から項目を構造化して取り出す。書式が揃っていなくても効く
  - ルール: 正規表現。APIキーが無い環境でも動く

どちらでも、確信が持てない件は書き込まず人へ引き継ぐ。
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from ...analysis.llm import ClaudeClient

# 台帳の列（実運用では既存ファイルの列に合わせる）
LEDGER_COLUMNS = ["受付日時", "氏名", "電話番号", "メールアドレス", "希望車種", "予算(円)", "メールID"]

# 人が見るべきメールのパターン
CLAIM_PATTERNS = re.compile(r"クレーム|苦情|至急|責任者|対応について|不具合|返金|解約")
MULTI_UNIT = re.compile(r"(?:合計\s*)?([0-9０-９]+)\s*台")

# ルール抽出用
NAME_PATTERN = re.compile(r"([一-龥ぁ-んァ-ヶーA-Za-z]{2,12})\s*(?:と申します|です[。\n])")
PHONE_PATTERN = re.compile(r"0\d{1,4}[-\s]?\d{1,4}[-\s]?\d{3,4}")
BUDGET_PATTERN = re.compile(r"([0-9０-９]+(?:\.[0-9]+)?)\s*万円")
KNOWN_MODELS = (
    "アクア", "フィット", "ノート", "ハスラー", "タント", "N-BOX", "セレナ",
    "ハリアー", "ヤリスクロス", "フォレスター", "CX-5", "RX", "ハイエース", "プリウス",
)

SYSTEM_PROMPT = """あなたは中古車販売店の事務担当です。
問い合わせメールから、顧客台帳に登録する項目を抜き出します。

厳守事項:
- 本文に書かれていないことを推測して埋めない。無い項目は null にする。
- 予算は円単位の整数にする（「130万円」→ 1300000）。
- 1通で複数台の相談をしている場合は multiple_vehicles を true にする。
- クレーム・苦情・至急の対応要請は claim を true にする。
- 出力は JSON オブジェクトのみ。説明を付けない。

形式:
{"name": "氏名 or null", "phone": "電話番号 or null", "model": "希望車種 or null",
 "budget_yen": 整数 or null, "multiple_vehicles": true/false, "claim": true/false,
 "confidence": 0.0〜1.0}"""


def _to_hankaku(text: str) -> str:
    return text.translate(str.maketrans("０１２３４５６７８９", "0123456789"))


class MailToLedgerRecipe:
    """問い合わせメール → 顧客台帳（Excel）。"""

    name = "mail_to_ledger"

    def __init__(
        self,
        mail_base_url: str,
        ledger: "LedgerWriter",
        client: ClaudeClient | None = None,
        timeout: float = 20.0,
    ) -> None:
        self.mail_base_url = mail_base_url.rstrip("/")
        self.ledger = ledger
        self.client = client
        self.timeout = timeout
        self._llm_failed = False

    # ------------------------------------------------------ 転記元(API)
    def fetch(self) -> Iterable[dict[str, Any]]:
        import httpx

        with httpx.Client(timeout=self.timeout) as client:
            res = client.get(
                f"{self.mail_base_url}/api/messages", params={"processed": "false"}
            )
            res.raise_for_status()
            return res.json()

    def key_of(self, source: dict[str, Any]) -> str:
        return str(source.get("id", "unknown"))

    # ------------------------------------------------------------ 判定
    def validate(self, source: dict[str, Any]) -> tuple[dict[str, Any] | None, str]:
        body = str(source.get("body") or "")
        fields, how = self._extract(body)

        if fields.get("claim"):
            return None, "クレーム・至急の対応要請の可能性があります。担当者が読んでください"
        if fields.get("multiple_vehicles"):
            return None, "複数台の相談のため、台帳の1行に収まりません"
        if not fields.get("model"):
            return None, f"希望車種を特定できませんでした（抽出: {how}）"
        if not fields.get("budget_yen"):
            return None, "予算が本文に書かれていません。確認が必要です"
        if not fields.get("name"):
            return None, "氏名を特定できませんでした"

        payload = {
            "受付日時": str(source.get("received_at") or "")[:16].replace("T", " "),
            "氏名": fields["name"],
            "電話番号": fields.get("phone") or "",
            "メールアドレス": str(source.get("from") or ""),
            "希望車種": fields["model"],
            "予算(円)": int(fields["budget_yen"]),
            "メールID": self.key_of(source),
        }
        return payload, f"{how}で全項目を抽出できました"

    def _extract(self, body: str) -> tuple[dict[str, Any], str]:
        if self.client and self.client.available and not self._llm_failed:
            result = self.client.complete_json(
                SYSTEM_PROMPT, f"次のメールから項目を抜き出してください。\n\n---\n{body}\n---",
                max_tokens=1000,
            )
            if isinstance(result, dict):
                return self._normalize(result), "AI抽出"
            # 一度失敗したらルールへ切り替える（毎件呼んで待たされないように）
            self._llm_failed = True
        return self._extract_by_rules(body), "ルール抽出"

    def _normalize(self, raw: dict[str, Any]) -> dict[str, Any]:
        budget = raw.get("budget_yen")
        try:
            budget = int(budget) if budget is not None else None
        except (TypeError, ValueError):
            budget = None
        return {
            "name": (raw.get("name") or None),
            "phone": (raw.get("phone") or None),
            "model": (raw.get("model") or None),
            "budget_yen": budget,
            "multiple_vehicles": bool(raw.get("multiple_vehicles")),
            "claim": bool(raw.get("claim")),
        }

    def _extract_by_rules(self, body: str) -> dict[str, Any]:
        text = _to_hankaku(body)
        name = NAME_PATTERN.search(body)
        phone = PHONE_PATTERN.search(text)
        budget = BUDGET_PATTERN.search(text)
        model = next((m for m in KNOWN_MODELS if m in body), None)

        units = [int(m) for m in MULTI_UNIT.findall(text) if m.isdigit()]
        return {
            "name": name.group(1) if name else None,
            "phone": phone.group(0) if phone else None,
            "model": model,
            "budget_yen": int(float(budget.group(1)) * 10000) if budget else None,
            "multiple_vehicles": any(u >= 2 for u in units),
            "claim": bool(CLAIM_PATTERNS.search(body)),
        }

    # -------------------------------------------------------- 転記先
    def apply(self, payload: dict[str, Any]) -> None:
        self.ledger.append_row(payload)

    def finalize(self, source: dict[str, Any], payload: dict[str, Any]) -> None:
        import httpx

        with httpx.Client(timeout=self.timeout) as client:
            res = client.post(
                f"{self.mail_base_url}/api/messages/{source['id']}/processed"
            )
            res.raise_for_status()


class LedgerWriter:
    """Excel の台帳へ1行追加する。

    既存ファイルがあれば追記し、無ければ見出し付きで作る。
    同じメールIDが既に入っていれば重複として拒否する
    （再実行しても二重登録にならないようにするため）。
    """

    def __init__(self, path: Path | str, sheet_name: str = "問い合わせ") -> None:
        self.path = Path(path)
        self.sheet_name = sheet_name

    @staticmethod
    def _is_empty(sheet) -> bool:
        return sheet.max_row <= 1 and all(c.value is None for c in sheet[1])

    def _open(self):
        from openpyxl import Workbook, load_workbook

        if self.path.exists():
            book = load_workbook(self.path)
            sheet = (
                book[self.sheet_name] if self.sheet_name in book.sheetnames
                else book.create_sheet(self.sheet_name)
            )
        else:
            book = Workbook()
            sheet = book.active
            sheet.title = self.sheet_name
        if self._is_empty(sheet):
            # 新規シートは「空だが1行ある」状態なので append だと2行目から書かれる。
            # 見出しは1行目へ直接置く。
            for column, name in enumerate(LEDGER_COLUMNS, start=1):
                sheet.cell(row=1, column=column, value=name)
        return book, sheet

    def existing_keys(self) -> set[str]:
        if not self.path.exists():
            return set()
        from openpyxl import load_workbook

        book = load_workbook(self.path, read_only=True)
        if self.sheet_name not in book.sheetnames:
            return set()
        sheet = book[self.sheet_name]
        index = LEDGER_COLUMNS.index("メールID")
        keys = set()
        for row in sheet.iter_rows(values_only=True):
            if not row or len(row) <= index:
                continue
            value = row[index]
            # 空行と見出し行は飛ばす
            if not value or str(value) == "メールID":
                continue
            keys.add(str(value))
        return keys

    def append_row(self, payload: dict[str, Any]) -> None:
        key = str(payload.get("メールID") or "")
        if key and key in self.existing_keys():
            raise RuntimeError(f"メールID {key} は既に台帳に登録済みです")
        self.path.parent.mkdir(parents=True, exist_ok=True)
        book, sheet = self._open()
        sheet.append([payload.get(col, "") for col in LEDGER_COLUMNS])
        book.save(self.path)

    def row_count(self) -> int:
        return len(self.existing_keys())


# ------------------------------------------------------------ 登録
def _build(spec, endpoints, ledger_path=None, client=None, **_):
    return MailToLedgerRecipe(
        mail_base_url=endpoints["source"],
        ledger=LedgerWriter(ledger_path or endpoints["target"]),
        client=client,
    )


def _register() -> None:
    from ..registry import RecipeEntry, register

    register(RecipeEntry(
        key="mail_to_ledger",
        label="メールの内容を台帳へ入力",
        categories=("データ入力",),
        needs_browser=False,
        description=(
            "受信メールから項目を抽出して台帳へ1行追加する。"
            "抽出はClaude（未設定時はルール）。確信が持てない件は書き込まず人へ回す。"
        ),
        factory=_build,
    ))


_register()
