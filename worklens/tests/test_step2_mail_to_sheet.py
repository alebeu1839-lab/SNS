"""メール→台帳レシピの検証。"""
from __future__ import annotations

import socket
import threading
import time

import httpx
import pytest
import uvicorn

from worklens.step2.recipes.mail_to_sheet import LedgerWriter, MailToLedgerRecipe
from worklens.step2.runner import AutomationRunner


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.fixture()
def mail_server():
    from mock.mail_system import MESSAGES, app

    for m in MESSAGES:
        m["processed"] = False
    port = _free_port()
    server = uvicorn.Server(
        uvicorn.Config(app, host="127.0.0.1", port=port, log_level="error")
    )
    threading.Thread(target=server.run, daemon=True).start()
    for _ in range(100):
        try:
            httpx.get(f"http://127.0.0.1:{port}/", timeout=0.5)
            break
        except Exception:
            time.sleep(0.05)
    yield f"http://127.0.0.1:{port}"
    server.should_exit = True
    for m in MESSAGES:
        m["processed"] = False


def _recipe(mail_url: str, tmp_path, client=None) -> MailToLedgerRecipe:
    return MailToLedgerRecipe(mail_url, LedgerWriter(tmp_path / "台帳.xlsx"), client=client)


# --------------------------------------------------------------- 抽出
def test_rule_extraction_reads_the_standard_inquiry(tmp_path):
    recipe = _recipe("http://x", tmp_path)
    payload, reason = recipe.validate({
        "id": "m1", "received_at": "2025-06-02T09:14:00+09:00", "from": "a@b.com",
        "body": "田中健一と申します。トヨタ アクアの購入を検討しています。\n"
                "予算は130万円程度です。090-1111-2222",
    })
    assert payload["氏名"] == "田中健一"
    assert payload["希望車種"] == "アクア"
    assert payload["予算(円)"] == 1300000
    assert payload["電話番号"] == "090-1111-2222"
    assert "ルール抽出" in reason


@pytest.mark.parametrize(
    "body, keyword",
    [
        ("高橋と申します。SUVを探しています。予算は決めていません。090-7777-8888", "希望車種"),
        ("伊藤です。至急、責任者の方からご連絡ください。アクアの件。予算は100万円。", "クレーム"),
        ("渡辺です。ハイエースを2台お願いします。予算は600万円です。", "複数台"),
        ("山田です。アクアが欲しいです。090-1111-2222", "予算"),
    ],
)
def test_uncertain_mails_go_to_a_human(tmp_path, body, keyword):
    recipe = _recipe("http://x", tmp_path)
    payload, reason = recipe.validate({"id": "m1", "received_at": "", "from": "a@b", "body": body})
    assert payload is None
    assert keyword in reason


class StubClaude:
    """LLM抽出の経路を、APIを呼ばずに確かめる。"""

    available = True
    model = "stub"

    def __init__(self, result):
        self.result = result
        self.calls = 0

    def complete_json(self, system, user, max_tokens=8000):
        self.calls += 1
        return self.result


def test_llm_extraction_is_used_when_available(tmp_path):
    client = StubClaude({
        "name": "山本花子", "phone": "090-0000-1111", "model": "ノート",
        "budget_yen": 990000, "multiple_vehicles": False, "claim": False,
    })
    recipe = _recipe("http://x", tmp_path, client=client)
    payload, reason = recipe.validate({
        "id": "m1", "received_at": "", "from": "a@b",
        "body": "書式のそろっていない自由文のメール本文",
    })
    assert client.calls == 1
    assert payload["氏名"] == "山本花子" and payload["希望車種"] == "ノート"
    assert "AI抽出" in reason


def test_falls_back_to_rules_when_llm_fails(tmp_path):
    client = StubClaude(None)          # APIが失敗した状況
    recipe = _recipe("http://x", tmp_path, client=client)
    mail = {"id": "m1", "received_at": "", "from": "a@b",
            "body": "佐藤美咲です。フィットで予算は150万円まで。080-3333-4444"}
    payload, reason = recipe.validate(mail)
    assert payload["氏名"] == "佐藤美咲"
    assert "ルール抽出" in reason
    # 一度失敗したら以降は呼ばない（毎件待たされないように）
    recipe.validate(mail)
    assert client.calls == 1


def test_llm_is_not_allowed_to_invent_missing_fields(tmp_path):
    """本文に無い項目をLLMが埋めても、欠けていれば人へ回す。"""
    client = StubClaude({
        "name": "推測太郎", "phone": None, "model": None, "budget_yen": None,
        "multiple_vehicles": False, "claim": False,
    })
    recipe = _recipe("http://x", tmp_path, client=client)
    payload, reason = recipe.validate({"id": "m1", "received_at": "", "from": "a@b",
                                       "body": "詳細未記載"})
    assert payload is None
    assert "希望車種" in reason


# ----------------------------------------------------------- 台帳への書き込み
def test_ledger_gets_a_header_row_once(tmp_path):
    writer = LedgerWriter(tmp_path / "台帳.xlsx")
    writer.append_row({"氏名": "田中", "メールID": "m1"})
    writer.append_row({"氏名": "佐藤", "メールID": "m2"})

    from openpyxl import load_workbook

    sheet = load_workbook(tmp_path / "台帳.xlsx")["問い合わせ"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "受付日時"        # 1行目が見出し（空行が入らない）
    assert sheet.max_row == 3
    assert writer.row_count() == 2


def test_ledger_refuses_duplicate_rows(tmp_path):
    writer = LedgerWriter(tmp_path / "台帳.xlsx")
    writer.append_row({"氏名": "田中", "メールID": "m1"})
    with pytest.raises(RuntimeError, match="既に台帳に登録済み"):
        writer.append_row({"氏名": "田中", "メールID": "m1"})


# ------------------------------------------------------------- 通し実行
def test_dry_run_writes_no_ledger_file(mail_server, tmp_path):
    recipe = _recipe(mail_server, tmp_path)
    result = AutomationRunner(recipe, tmp_path, mode="dry-run").run()
    assert result.processed == 7
    assert result.succeeded == 4
    assert result.handoff == 3
    assert not (tmp_path / "台帳.xlsx").exists()
    assert all(not m["processed"] for m in httpx.get(f"{mail_server}/api/messages").json())


def test_live_run_writes_only_the_certain_ones(mail_server, tmp_path):
    recipe = _recipe(mail_server, tmp_path)
    result = AutomationRunner(recipe, tmp_path, mode="live").run()
    assert result.succeeded == 4 and result.handoff == 3 and result.failed == 0

    assert LedgerWriter(tmp_path / "台帳.xlsx").existing_keys() == {"m001", "m002", "m003", "m007"}
    unprocessed = [
        m["id"] for m in httpx.get(f"{mail_server}/api/messages").json() if not m["processed"]
    ]
    assert sorted(unprocessed) == ["m004", "m005", "m006"]


def test_rerun_does_not_duplicate_ledger_rows(mail_server, tmp_path):
    recipe = _recipe(mail_server, tmp_path)
    AutomationRunner(recipe, tmp_path, mode="live").run()
    second = AutomationRunner(recipe, tmp_path, mode="live").run()
    assert second.processed == 3 and second.succeeded == 0
    assert LedgerWriter(tmp_path / "台帳.xlsx").row_count() == 4
